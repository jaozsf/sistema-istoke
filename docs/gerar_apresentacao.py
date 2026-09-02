"""
Gera o slide da Etapa 1 (Apresentação Inicial da Empresa) em .pptx.
Uso: python gerar_apresentacao.py
Saída: Trabalho_Etapa1_StockIQ.pptx (mesma pasta)
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

PRIMARY = RGBColor(0x53, 0x4A, 0xB7)
DARK    = RGBColor(0x1A, 0x1A, 0x19)
MUTED   = RGBColor(0x66, 0x66, 0x64)
LIGHT   = RGBColor(0xF5, 0xF5, 0xF0)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
SUCCESS = RGBColor(0x1D, 0x9E, 0x75)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def bg(slide, color=WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def bar(slide, color=PRIMARY, height=Inches(0.14)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def title(slide, text, top=Inches(0.5), color=DARK, size=36):
    box = slide.shapes.add_textbox(Inches(0.7), top, Inches(12), Inches(1))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = color
    return box


def subtitle(slide, text, top, color=MUTED, size=18, left=Inches(0.7), width=Inches(11.9)):
    box = slide.shapes.add_textbox(left, top, width, Inches(0.6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    return box


def bullets(slide, items, top, left=Inches(0.8), width=Inches(11.6), size=18, gap=Pt(10)):
    box = slide.shapes.add_textbox(left, top, width, Inches(5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            head, rest = item
            r1 = p.add_run(); r1.text = f"{head}  "
            r1.font.bold = True; r1.font.size = Pt(size); r1.font.color.rgb = PRIMARY
            r2 = p.add_run(); r2.text = rest
            r2.font.size = Pt(size); r2.font.color.rgb = DARK
        else:
            p.text = f"•  {item}"
            p.font.size = Pt(size)
            p.font.color.rgb = DARK
        p.space_after = gap
    return box


def footer(slide, text="StockIQ · Etapa 1 — Apresentação Inicial"):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(7.05), Inches(11.9), Inches(0.35))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(11)
    p.font.color.rgb = MUTED


# ── Slide 1 — Capa ───────────────────────────────────────────────────────────
s = add_slide()
bg(s, PRIMARY)
logo = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.67), Inches(1.9), Inches(2), Inches(2))
logo.fill.solid(); logo.fill.fore_color.rgb = WHITE; logo.line.fill.background()
tf = logo.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.text = "IQ"; p.font.size = Pt(60); p.font.bold = True
p.font.color.rgb = PRIMARY; p.alignment = PP_ALIGN.CENTER
tf.vertical_anchor = MSO_ANCHOR.MIDDLE

box = s.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.33), Inches(1.2))
p = box.text_frame.paragraphs[0]
p.text = "StockIQ"
p.font.size = Pt(54); p.font.bold = True; p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER

box = s.shapes.add_textbox(Inches(1), Inches(5.15), Inches(11.33), Inches(0.6))
p = box.text_frame.paragraphs[0]
p.text = "ERP inteligente de estoque para pequenos negócios"
p.font.size = Pt(20); p.font.color.rgb = RGBColor(0xE5, 0xE2, 0xF5)
p.alignment = PP_ALIGN.CENTER

box = s.shapes.add_textbox(Inches(1), Inches(6.3), Inches(11.33), Inches(0.6))
p = box.text_frame.paragraphs[0]
p.text = "Etapa 1 — Apresentação Inicial da Empresa"
p.font.size = Pt(15); p.font.color.rgb = RGBColor(0xC9, 0xC4, 0xE8)
p.alignment = PP_ALIGN.CENTER

# ── Slide 2 — Identidade ─────────────────────────────────────────────────────
s = add_slide(); bg(s); bar(s)
title(s, "Quem somos")
subtitle(s, "Nome fantasia: StockIQ", Inches(1.35), color=DARK, size=22)
bullets(s, [
    ("Segmento:", "software (ERP SaaS) para gestão de estoque"),
    ("Fundação:", "2026, como projeto acadêmico transformado em produto real"),
    ("Identidade visual:", "roxo #534AB7 — tecnologia, confiança e clareza"),
], Inches(2.3))
footer(s)

# ── Slide 3 — Missão, Visão, Valores ─────────────────────────────────────────
s = add_slide(); bg(s); bar(s)
title(s, "Missão, Visão e Valores")
bullets(s, [
    ("Missão:", "dar a pequenos e médios negócios o mesmo controle de estoque que só grandes empresas têm hoje — simples, rápido e sem planilha."),
    ("Visão:", "ser o ERP padrão de quem vende produto físico e cansou de perder dinheiro sem saber onde."),
    ("Valores:", "simplicidade antes de feature nova · dado real antes de achismo · preço justo para quem tá começando."),
], Inches(1.6), size=17)
footer(s)

# ── Slide 4 — Equipe ──────────────────────────────────────────────────────────
s = add_slide(); bg(s); bar(s)
title(s, "Equipe")
team = [
    ("Matheus", "CEO"),
    ("João Bazarin", "CTO"),
    ("Otávio", "Product Manager"),
    ("Giovani", "Dev / QA"),
]
card_w = Inches(2.7); gap = Inches(0.35)
start_x = Inches(0.85)
for i, (name, role) in enumerate(team):
    x = start_x + i * (card_w + gap)
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.3), card_w, Inches(2.6))
    card.fill.solid(); card.fill.fore_color.rgb = LIGHT
    card.line.color.rgb = RGBColor(0xE8, 0xE6, 0xDE); card.line.width = Pt(1)
    avatar = s.shapes.add_shape(MSO_SHAPE.OVAL, x + card_w/2 - Inches(0.5), Inches(2.6), Inches(1), Inches(1))
    avatar.fill.solid(); avatar.fill.fore_color.rgb = PRIMARY; avatar.line.fill.background()
    atf = avatar.text_frame; atf.vertical_anchor = MSO_ANCHOR.MIDDLE
    ap = atf.paragraphs[0]; ap.text = name[0]; ap.font.size = Pt(28); ap.font.bold = True
    ap.font.color.rgb = WHITE; ap.alignment = PP_ALIGN.CENTER

    ntf = s.shapes.add_textbox(x, Inches(3.75), card_w, Inches(0.5)).text_frame
    npar = ntf.paragraphs[0]; npar.text = name; npar.font.size = Pt(18); npar.font.bold = True
    npar.font.color.rgb = DARK; npar.alignment = PP_ALIGN.CENTER

    rtf = s.shapes.add_textbox(x, Inches(4.2), card_w, Inches(0.5)).text_frame
    rpar = rtf.paragraphs[0]; rpar.text = role; rpar.font.size = Pt(15)
    rpar.font.color.rgb = PRIMARY; rpar.alignment = PP_ALIGN.CENTER
footer(s)

# ── Slide 5 — Mercado-alvo ────────────────────────────────────────────────────
s = add_slide(); bg(s); bar(s)
title(s, "Área de atuação e mercado-alvo")
bullets(s, [
    ("Segmento:", "software de gestão (ERP) — módulo de estoque, financeiro básico e IA"),
    ("Cliente-alvo:", "pequenos negócios com produto físico em caixa/lote — distribuidoras, lojas, indústrias de pequeno porte"),
    ("Caso piloto real:", "distribuidora de copos e canecas acrílicas, controle por caixa fechada de 50 unidades"),
    ("Tamanho do problema:", "maioria desses negócios ainda controla estoque em planilha ou caderno"),
], Inches(1.6), size=18)
footer(s)

# ── Slide 6 — Ideia do produto ────────────────────────────────────────────────
s = add_slide(); bg(s); bar(s)
title(s, "Ideia inicial do produto")
subtitle(s, "StockIQ — ERP com controle de estoque via QR Code", Inches(1.35), color=DARK, size=20)
bullets(s, [
    "App (web + mobile) onde cada produto tem um QR Code próprio",
    "Escanear a caixa já registra a movimentação de estoque na hora",
    "Dashboard mostra saldo real, receita e lucro do dia — sem planilha",
    "Assistente respondese perguntas do negócio ('onde tô perdendo dinheiro?') na hora",
    "Login simples: e-mail/senha ou Google, acesso controlado por permissão (admin/gerente/operador)",
], Inches(2.3))
footer(s)

# ── Slide 7 — Problema ────────────────────────────────────────────────────────
s = add_slide(); bg(s); bar(s)
title(s, "Problema a ser resolvido")
bullets(s, [
    "Contagem de estoque manual: demorada e sujeita a erro humano",
    "Sem visibilidade de margem real por produto — preço definido \"no olho\"",
    "Estoque baixo só é percebido quando o produto já acabou",
    "Nenhum registro de quem retirou o quê, quando e por quê",
    "Ferramentas de ERP grandes são caras e complexas demais pra esse porte de negócio",
], Inches(1.6))
footer(s)

# ── Slide 8 — Stack ───────────────────────────────────────────────────────────
s = add_slide(); bg(s); bar(s)
title(s, "Tecnologias previstas (stack)")
cols = [
    ("Backend", ["FastAPI + Python", "PostgreSQL + SQLAlchemy (async)", "Alembic (migrações)", "Docker + Docker Compose"]),
    ("Frontend", ["Flutter (web, Android, iOS)", "Riverpod (estado)", "go_router (navegação)"]),
    ("Plataforma", ["Firebase Auth (e-mail + Google)", "JWT interno", "QR Code (geração + leitura)", "IA: respostas padrão + Claude (Anthropic) opcional"]),
]
col_w = Inches(3.9); gap = Inches(0.25); start_x = Inches(0.7)
for i, (head, items) in enumerate(cols):
    x = start_x + i * (col_w + gap)
    htf = s.shapes.add_textbox(x, Inches(1.6), col_w, Inches(0.5)).text_frame
    hp = htf.paragraphs[0]; hp.text = head; hp.font.size = Pt(19); hp.font.bold = True; hp.font.color.rgb = PRIMARY
    btf = s.shapes.add_textbox(x, Inches(2.15), col_w, Inches(4)).text_frame
    btf.word_wrap = True
    for j, item in enumerate(items):
        p = btf.paragraphs[0] if j == 0 else btf.add_paragraph()
        p.text = f"•  {item}"; p.font.size = Pt(15); p.font.color.rgb = DARK
        p.space_after = Pt(8)
footer(s)

# ── Slide 9 — Diferencial ──────────────────────────────────────────────────────
s = add_slide(); bg(s); bar(s)
title(s, "Diferencial competitivo")
bullets(s, [
    ("Movimentação por QR Code:", "escaneou, já baixou/subiu o estoque — sem digitar nada"),
    ("IA sem custo pra maioria das perguntas:", "padrões comuns (estoque baixo, produto parado, prejuízo) respondem na hora com dado real, sem gastar API de IA"),
    ("Multi-plataforma desde o dia 1:", "mesmo código roda em navegador, Android e iOS"),
    ("Preço acessível:", "pensado pra pequeno negócio, não pra corporação"),
], Inches(1.6), size=18)
footer(s)

# ── Slide 10 — Cronograma ─────────────────────────────────────────────────────
s = add_slide(); bg(s); bar(s)
title(s, "Cronograma macro")
stages = [
    ("Etapa 1", "Apresentação da empresa e ideia (agora)"),
    ("Etapa 2", "MVP funcional: cadastro de produto, estoque, QR Code, login"),
    ("Etapa 3", "Financeiro, dashboard com dados reais, testes com usuário piloto"),
    ("Etapa 4", "IA, relatórios em PDF, ajustes finais e entrega"),
]
top = Inches(1.9)
for i, (stage, desc) in enumerate(stages):
    y = top + i * Inches(1.15)
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), y, Inches(0.35), Inches(0.35))
    dot.fill.solid(); dot.fill.fore_color.rgb = PRIMARY if i > 0 else SUCCESS; dot.line.fill.background()
    stf = s.shapes.add_textbox(Inches(1.4), y - Inches(0.08), Inches(2), Inches(0.5)).text_frame
    sp = stf.paragraphs[0]; sp.text = stage; sp.font.bold = True; sp.font.size = Pt(18); sp.font.color.rgb = DARK
    dtf = s.shapes.add_textbox(Inches(3.3), y - Inches(0.08), Inches(9), Inches(0.5)).text_frame
    dp = dtf.paragraphs[0]; dp.text = desc; dp.font.size = Pt(16); dp.font.color.rgb = MUTED
    if i < len(stages) - 1:
        line = s.shapes.add_connector(1, Inches(0.975), y + Inches(0.35), Inches(0.975), y + Inches(1.15))
        line.line.color.rgb = RGBColor(0xD8, 0xD5, 0xCC); line.line.width = Pt(2)
footer(s)

# ── Slide 11 — Encerramento ───────────────────────────────────────────────────
s = add_slide(); bg(s, PRIMARY)
box = s.shapes.add_textbox(Inches(1), Inches(3), Inches(11.33), Inches(1))
p = box.text_frame.paragraphs[0]
p.text = "Obrigado!"
p.font.size = Pt(48); p.font.bold = True; p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER
box = s.shapes.add_textbox(Inches(1), Inches(4), Inches(11.33), Inches(0.6))
p = box.text_frame.paragraphs[0]
p.text = "Perguntas? Matheus (CEO) · João Bazarin (CTO) · Otávio (PM) · Giovani (Dev/QA)"
p.font.size = Pt(16); p.font.color.rgb = RGBColor(0xE5, 0xE2, 0xF5)
p.alignment = PP_ALIGN.CENTER

prs.save("Trabalho_Etapa1_StockIQ.pptx")
print("Gerado: Trabalho_Etapa1_StockIQ.pptx")
